#!/usr/bin/env python3
"""Rebuild the FTL violation log.

Scans every .xlsx export under ~/Downloads/FTL_violations (any subfolder),
merges them, dedupes exact-duplicate rows, then collapses re-exports of the
same (pilot, date, route) down to the row from the newest file (recalculation
drift across daily exports otherwise creates near-duplicate rows). Computes
why each row was flagged, and links duty pairs flown by two pilots so a
single "handled" mark covers both. Writes data/final_records.json and
regenerates index.html from template.html.

Usage: ./venv/bin/python build.py   (or just ./rebuild.sh, which also commits+pushes)
"""
import glob
import os
import re
import json
import hashlib
import datetime
from collections import defaultdict

import openpyxl

SRC_DIR = os.path.expanduser("~/Downloads/FTL_violations")
REPO_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_FILE = os.path.join(REPO_DIR, "data", "final_records.json")
SEED_HANDLED_FILE = os.path.join(REPO_DIR, "data", "seed_handled.json")
TEMPLATE_FILE = os.path.join(REPO_DIR, "template.html")
OUTPUT_FILE = os.path.join(REPO_DIR, "index.html")

SHARED_DUTY_FIELDS = ["Duty end date [UTC]", "Route ICAO", "Duty end time [UTC]", "FDP end time [UTC]"]
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def norm(v):
    if isinstance(v, (datetime.datetime, datetime.time)):
        return v.isoformat()
    return "" if v is None else str(v)


def minutes(s):
    if not s:
        return None
    if s.startswith(">"):
        return 999999
    parts = s.split(":")
    try:
        h, m = int(parts[0]), int(parts[1])
        sec = int(parts[2]) if len(parts) == 3 else 0
        return h * 60 + m + sec / 60
    except ValueError:
        return None


def load_previous_ids():
    if not os.path.exists(DATA_FILE):
        return set()
    with open(DATA_FILE) as fh:
        return {r["_id"] for r in json.load(fh)}


def read_rows():
    files = sorted(glob.glob(f"{SRC_DIR}/**/*.xlsx", recursive=True))
    if not files:
        raise SystemExit(f"No .xlsx files found under {SRC_DIR}")
    headers = None
    rows = {}  # full-row key -> (record dict, newest mtime seen for this exact row)
    n_raw = 0
    for f in files:
        ws = openpyxl.load_workbook(f, data_only=True)["ftl_duty"]
        h = [c.value for c in ws[1]]
        headers = headers or h
        assert h == headers, f"Unexpected column layout in {f}"
        mtime = os.path.getmtime(f)
        for row in ws.iter_rows(min_row=2, values_only=True):
            if all(c is None for c in row):
                continue
            n_raw += 1
            key = tuple(norm(v) for v in row)
            if key in rows:
                rows[key] = (rows[key][0], max(rows[key][1], mtime))
                continue
            rows[key] = (dict(zip(headers, key)), mtime)
    return files, rows, n_raw


def compute_reasons(r):
    reasons = []
    fdp, mx = minutes(r["FDP length"]), minutes(r["Max FDP"])
    if fdp is not None and mx is not None and fdp > mx:
        reasons.append(f"FDP length {r['FDP length']} exceeds Max FDP {r['Max FDP']}")

    # Only rest-BEFORE counts as a primary reason: the rest gap after duty N is the
    # same physical gap as the rest gap before duty N+1, so checking both sides
    # would flag one shortfall twice across two different rows.
    req_b, act_b = minutes(r["Rest required before"]), minutes(r["Rest actual before"])
    if req_b is not None and act_b is not None and act_b < req_b:
        reasons.append(f"Rest before ({r['Rest actual before']}) below required ({r['Rest required before']})")

    d, p = minutes(r["Duty length"]), minutes(r["Planned duty length"])
    if d is not None and p is not None and d > p + 1:
        reasons.append(f"Duty length {r['Duty length']} exceeds planned {r['Planned duty length']}")

    if not reasons:
        req_a, act_a = minutes(r["Rest required after"]), minutes(r["Rest actual after"])
        if req_a is not None and act_a is not None and act_a < req_a:
            reasons.append(
                f"Rest after this duty ({r['Rest actual after']}) below required ({r['Rest required after']}) "
                f"— same gap is flagged as 'Rest before' on the next duty"
            )
    return reasons or ["Flagged by system (no threshold auto-detected — review manually)"]


def assign_links(records):
    groups = defaultdict(list)
    for r in records:
        groups[tuple(r[f] for f in SHARED_DUTY_FIELDS)].append(r)
    for key, g in groups.items():
        lid = hashlib.sha1(json.dumps(key).encode()).hexdigest()[:12]
        for r in g:
            r["_link_id"] = lid if len(g) > 1 else r["_id"]
            r["_link_partners"] = [x["Crew member"] for x in g if x is not r]
    return sum(len(g) > 1 for g in groups.values())


def build_records():
    files, rows, n_raw = read_rows()
    previous_ids = load_previous_ids()

    records = []
    for r, mtime in rows.values():
        r["Crew member"] = re.sub(r"\s+", " ", r["Crew member"]).strip()
        r["Duty end date [UTC]"] = r["Duty end date [UTC]"].split("T")[0]
        r["_reasons"] = compute_reasons(r)
        r["_id"] = hashlib.sha1(
            "|".join([r["Crew member"], r["Crew code"], r["Duty end date [UTC]"], r["Route ICAO"],
                      r["Duty end time [UTC]"], r["FDP end time [UTC]"]]).encode()
        ).hexdigest()[:12]
        r["_export_mtime"] = mtime
        records.append(r)

    records.sort(key=lambda r: (r["Crew member"], r["Duty end date [UTC]"], r["Duty end time [UTC]"]))
    assign_links(records)
    old_link = {r["_id"]: r["_link_id"] for r in records}

    # Collapse re-exports of the same (pilot, date, route) to the row from the newest file.
    dup = defaultdict(list)
    for r in records:
        dup[(r["Crew member"], r["Duty end date [UTC]"], r["Route ICAO"])].append(r)
    kept, dropped = [], []
    for key, g in dup.items():
        if len(g) == 1:
            kept.append(g[0])
            continue
        newest = max(x["_export_mtime"] for x in g)
        top = [x for x in g if x["_export_mtime"] == newest]
        if len(top) > 1:
            kept.extend(g)  # tie: don't guess, keep all
            continue
        win = top[0]
        win["_aliases"] = sorted({old_link[x["_id"]] for x in g} - {win["_id"]})
        kept.append(win)
        dropped.extend(x for x in g if x is not win)

    records = sorted(kept, key=lambda r: (r["Crew member"], r["Duty end date [UTC]"], r["Duty end time [UTC]"]))
    for r in records:
        r.setdefault("_aliases", [])
    n_pairs = assign_links(records)
    for r in records:
        r["_aliases"] = [a for a in r["_aliases"] if a != r["_link_id"]]
        r["_new"] = r["_id"] not in previous_ids
        del r["_export_mtime"]

    return records, files, n_raw, dropped, n_pairs


def main():
    records, files, n_raw, dropped, n_pairs = build_records()

    os.makedirs(os.path.dirname(DATA_FILE), exist_ok=True)
    with open(DATA_FILE, "w") as fh:
        json.dump(records, fh, indent=2, ensure_ascii=False)

    latest = max(r["Duty end date [UTC]"] for r in records)
    y, m, d = latest.split("-")
    seed_handled = {}
    if os.path.exists(SEED_HANDLED_FILE):
        with open(SEED_HANDLED_FILE) as fh:
            seed_handled = json.load(fh)
    with open(TEMPLATE_FILE) as fh:
        html = fh.read()
    html = (
        html.replace("__RECORDS_JSON__", json.dumps(records, ensure_ascii=False))
        .replace("__SEED_HANDLED_JSON__", json.dumps(seed_handled, ensure_ascii=False))
        .replace("__FILE_COUNT__", str(len(files)))
        .replace("__LATEST_DATE__", f"{int(d)} {MONTHS[int(m) - 1]}")
        .replace("__BUILD_DATE__", datetime.date.today().isoformat())
    )
    with open(OUTPUT_FILE, "w") as fh:
        fh.write(html)

    print(
        f"{len(files)} xlsx files | {n_raw} raw rows | {len(records) + len(dropped)} unique | "
        f"{len(records)} after collapse ({len(dropped)} superseded) | "
        f"{len({r['Crew member'] for r in records})} pilots | {n_pairs} paired duties | "
        f"{sum(r['_new'] for r in records)} new since last build"
    )


if __name__ == "__main__":
    main()
